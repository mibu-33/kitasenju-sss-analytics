"""Slow, read-only collection of publicly available SloRepo data.

Missing observations remain missing. Never bypass site controls or fabricate rows.
"""
import datetime as dt
import io
from itertools import chain
import json
import os
import re
import time
from pathlib import Path
from urllib.parse import urljoin
from zoneinfo import ZoneInfo

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

BASE = 'https://www.slorepo.com/hole/e38394e383bce382a2e383bce382afe58c97e58d83e4bd8f535353code/'
ROOT = Path(__file__).resolve().parents[1] / 'data'
HEADERS = {'User-Agent': 'SSSAnalytics/1.1 (+personal research; respectful low-frequency fetching)'}
SESSION = requests.Session()
SESSION.headers.update(HEADERS)
TIMEZONE = ZoneInfo('Asia/Tokyo')


def number(value):
    s = re.sub(r'[,，\s枚ＧG%＋+]', '', str(value or '')).replace('−','-').replace('ー','-')
    m = re.search(r'-?\d+(?:\.\d+)?', s)
    return float(m.group()) if m else None


def as_date(value, year=None):
    if isinstance(value, dt.datetime): return value.date().isoformat()
    if isinstance(value, dt.date): return value.isoformat()
    s = str(value or '').strip().replace('年','/').replace('月','/').replace('日','')
    for pat in (r'(20\d\d)[/-](\d{1,2})[/-](\d{1,2})', r'^(20\d\d)(\d\d)(\d\d)$'):
        match = re.search(pat, s)
        if match:
            try: return dt.date(*map(int, match.groups())).isoformat()
            except ValueError: return None
    if year is not None:
        match = re.search(r'^(\d{1,2})[/-](\d{1,2})(?:\D|$)', s)
        if match:
            try: return dt.date(year, *map(int, match.groups())).isoformat()
            except ValueError: return None
    return None


def fetch(url, binary=False):
    if not url.startswith(BASE): raise ValueError('取得元が指定店舗のURLではありません')
    r = SESSION.get(url, timeout=28, allow_redirects=True)
    r.raise_for_status()
    if not r.url.startswith(BASE): raise ValueError('別サイトへ転送されたため中断')
    if len(r.content) > 18_000_000: raise ValueError('ファイル容量が大きすぎます')
    if any(s in r.text[:1500].lower() for s in ['captcha', 'access denied','cloudflare challenge']) if not binary else False:
        raise ValueError('サイト側のアクセス制限により取得を停止')
    if binary: return r.content
    r.encoding = r.apparent_encoding or 'utf-8'
    return BeautifulSoup(r.text, 'html.parser')


def table_rows(table):
    return [[cell.get_text(' ',strip=True) for cell in row.find_all(['td','th'])] for row in table.select('tr') if row.find_all(['td','th'])]


def column(header, names):
    normalized=[re.sub(r'[\s　]', '', str(h or '')).lower() for h in header]
    return next((i for i,v in enumerate(normalized) if any(name in v for name in names)), None)


def win_parts(raw):
    m=re.search(r'(\d+)\s*/\s*(\d+)',str(raw or ''))
    return (int(m.group(1)),int(m.group(2))) if m else (None,None)


def summary_from_index(soup, today=None):
    today = today or dt.datetime.now(TIMEZONE).date()
    out={}; links={}
    for anchor in soup.find_all('a',href=True):
        match=re.search(r'(?<!\d)(\d{1,2})/(\d{1,2})\s*\(',anchor.get_text(' ',strip=True))
        if not match: continue
        month,day=map(int,match.groups())
        try:
            value=dt.date(today.year,month,day)
            if value>today: value=value.replace(year=today.year-1)
        except ValueError: continue
        url=urljoin(BASE,anchor['href']).rstrip('/')+'/'
        if not url.startswith(BASE) or not re.search(r'/\d{8}/$',url): continue
        date=value.isoformat()
        # Use URL's 8-digit date as authority for the year, not date in the user's OS.
        url_date=as_date(url.rstrip('/').split('/')[-1])
        if url_date: date=url_date
        links[date]=url
        row=anchor.find_parent('tr')
        if not row: continue
        vals=[c.get_text(' ',strip=True) for c in row.find_all('td')]
        if len(vals)<4: continue
        avg, games = number(vals[1]), number(vals[2])
        if avg is None or games is None: continue
        winners,count=win_parts(vals[3]); pct=round(100*winners/count,2) if count else number(vals[3])
        out[date]={'date':date,'avg_diff':avg,'avg_games':games,'win_rate':pct,'wins':winners,'count':count,'source':'slorepo-index'}
    return list(out.values()),links


def extract_day(soup, date):
    """If per-machine tables are available, import rows only with explicit machine number and games and diff."""
    result=[]
    for table in soup.find_all('table'):
        rows=table_rows(table)
        for idx,headers in enumerate(rows[:4]):
            n=column(headers,['台番','台番号']); diff=column(headers,['差枚','差玉']); games=column(headers,['g数','ゲーム数','回転数','総回転']); machine=column(headers,['機種']);bb=column(headers,['bb']);rb=column(headers,['rb'])
            if None in (n,diff,games): continue
            for cells in rows[idx+1:]:
                if max(n,diff,games)>=len(cells): continue
                raw=str(cells[n]).strip().replace(',','')
                if not re.fullmatch(r'\d{1,5}',raw): continue
                d=number(cells[diff]);g=number(cells[games])
                if d is None or g is None or not (-30000<=d<=30000) or not (0<=g<=50000): continue
                result.append({'date':date,'number':int(raw),'machine':cells[machine] if machine is not None and machine<len(cells) else '機種不明','diff':int(d),'games':int(g),'bb':int(number(cells[bb])) if bb is not None and bb<len(cells) and number(cells[bb]) is not None else None,'rb':int(number(cells[rb])) if rb is not None and rb<len(cells) and number(cells[rb]) is not None else None})
            break
    return list({(x['date'],x['number']):x for x in result}.values())


def extract_aggregate(soup,date):
    machines=[];tails=[]
    for table in soup.find_all('table'):
        rows=table_rows(table)
        if len(rows)<2:continue
        for idx,headers in enumerate(rows[:3]):
            name=column(headers,['機種','末尾番号']);diff=column(headers,['平均差枚']);games=column(headers,['平均g','平均ゲーム']);win=column(headers,['勝率'])
            if None in (name,diff,games,win): continue
            is_tail=any('末尾番号' in str(h) for h in headers)
            for cells in rows[idx+1:]:
                if max(name,diff,games,win)>=len(cells): continue
                label=cells[name].strip(); d=number(cells[diff]);g=number(cells[games]);w,n=win_parts(cells[win])
                if not label or d is None or g is None or n is None:continue
                item={'date':date,'avg_diff':d,'avg_games':g,'wins':w,'count':n,'win_rate':round(100*w/n,2) if n else None}
                if is_tail:
                    item['tail']=label.replace('末尾','');tails.append(item)
                else:
                    item['machine']=label;machines.append(item)
            break
    return machines,tails


def normalize_col(value):
    return re.sub(r'[\s　_()（）\-/]', '', str(value or '')).lower()


def parse_raw_xlsx(content, sheet_fallback_date=None):
    """Workbook ingestion only when each row has a verifiable date or dated sheet.
    This intentionally rejects aggregate-only sheets without seat numbers.
    """
    if not content.startswith(b'PK'):
        raise ValueError('Excelファイルではありません')
    wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    results=[]
    for ws in wb.worksheets:
        inferred=as_date(ws.title) or sheet_fallback_date
        iterator=ws.iter_rows(values_only=True)
        window=[]
        for _ in range(25):
            try:window.append(next(iterator))
            except StopIteration:break
        idxs=None;header_index=None
        for i,row in enumerate(window):
            h=[normalize_col(c) for c in row]
            def col(*xs): return next((n for n,cell in enumerate(h) if any(x in cell for x in xs)),None)
            seat=col('台番号','台番'); diff=col('差枚','差玉');games=col('g数','ゲーム数','回転数','総回転')
            if None in (seat,diff,games):continue
            idxs=(seat,diff,games,col('機種'),col('日付','年月日','date'),col('bb'),col('rb'));header_index=i;break
        if idxs is None:continue
        rows=chain(window[header_index+1:],iterator)
        seat,diff,games,machine,date,bb,rb=idxs
        for row in rows:
            if not row or max(seat,diff,games)>=len(row):continue
            d=as_date(row[date]) if date is not None and date<len(row) else inferred
            if d is None:continue
            n=number(row[seat]);delta=number(row[diff]);g=number(row[games])
            if n is None or n!=int(n) or not (1<=n<=99999) or delta is None or g is None:continue
            if not (-30000<=delta<=30000) or not (0<=g<=50000):continue
            def optional_int(i):
                v=number(row[i]) if i is not None and i<len(row) else None
                return int(v) if v is not None else None
            results.append({'date':d,'number':int(n),'machine':str(row[machine] or '機種不明').strip() if machine is not None and machine<len(row) else '機種不明','diff':int(delta),'games':int(g),'bb':optional_int(bb),'rb':optional_int(rb)})
    return list({(x['date'],x['number']):x for x in results}.values())


def read(name,default):
    try:return json.loads((ROOT/name).read_text(encoding='utf-8'))
    except (FileNotFoundError,ValueError):return default


def save(name,value):
    (ROOT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def main():
    ROOT.mkdir(exist_ok=True)
    now=dt.datetime.now(TIMEZONE).isoformat(timespec='seconds')
    existing={(x['date'],x['number']):x for x in read('records.json',[])}
    summaries={x['date']:x for x in read('summary.json',[])}
    machine_days={(x['date'],x['machine']):x for x in read('machine_daily.json',[])}
    tail_days={(x['date'],x['tail']):x for x in read('tail_daily.json',[])}
    status=read('status.json',{})
    status['last_attempt']=now;errors=[];rows_imported=0;agg_imported=0
    try:
        soup=fetch(BASE);indexed,links=summary_from_index(soup)
        if not links: raise ValueError('有効な日別レポートリンクを発見できません')
        for s in indexed:summaries[s['date']]=s
        # RAW XLSX is the preferred data source for machine-number-level data.
        try:
            raw=fetch(urljoin(BASE,'table_data.xlsx'),binary=True)
            parsed=parse_raw_xlsx(raw)
            for row in parsed:existing[(row['date'],row['number'])]=row
            rows_imported+=len(parsed)
            if not parsed:errors.append('RAW Excelに日付・台番号・差枚・G数が揃う表がありません')
        except Exception as e:
            errors.append(f'RAW Excel: {type(e).__name__}: {e}')
        days=max(1,min(int(os.getenv('SYNC_DAYS','7')),20))
        candidates=sorted(links,reverse=True)
        fresh=candidates[:min(3,days)]
        # Limit requests and keep historical data; older unindexed days are never invented.
        pending=[d for d in candidates if not any(key[0]==d for key in machine_days)]
        targets=list(dict.fromkeys(fresh+pending))[:days]
        for d in targets:
            try:
                doc=fetch(links[d]);m,t=extract_aggregate(doc,d)
                for x in m:machine_days[(x['date'],x['machine'])]=x
                for x in t:tail_days[(x['date'],x['tail'])]=x
                agg_imported+=len(m)+len(t)
                rows=extract_day(doc,d)
                for x in rows:existing[(x['date'],x['number'])]=x
                rows_imported+=len(rows)
                if not m and not t and not rows:errors.append(f'{d}: 集計テーブルなし')
                time.sleep(max(1.5,float(os.getenv('REQUEST_INTERVAL','2.5'))))
            except Exception as e:errors.append(f'{d}: {type(e).__name__}: {e}')
        status['last_success']=now
        status['newest_date']=max(summaries,default=None)
        status['status']='partial' if errors else 'ok'
    except Exception as e:
        errors.append(f'店舗ページ: {type(e).__name__}: {e}')
        status['status']='error'
    status.update(error='; '.join(errors)[:1600] if errors else None,rows_imported=rows_imported,aggregates_imported=agg_imported,dated_machine_rows=len(existing),indexed_days=len(summaries),machine_daily_rows=len(machine_days))
    save('records.json',sorted(existing.values(),key=lambda x:(x['date'],x['number'])))
    save('summary.json',sorted(summaries.values(),key=lambda x:x['date']))
    save('machine_daily.json',sorted(machine_days.values(),key=lambda x:(x['date'],x['machine'])))
    save('tail_daily.json',sorted(tail_days.values(),key=lambda x:(x['date'],x['tail'])))
    save('status.json',status)
    print(json.dumps(status,ensure_ascii=False))
    if status['status']=='error':raise SystemExit(1)

if __name__=='__main__':main()
